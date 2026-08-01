import os
import re
from datetime import date, datetime
from openpyxl import Workbook, load_workbook

EXCEL_DIR = "excel_files"

COLUMNS = [
    "S_No", "Date", "Roll_No", "Register_No", "Student_Name", "Father_Name", "Class", "Course", "Academic_Year",
    "Tuition_Fee_Year1", "Bus_Fee_Year1", "Total_Year1",
    "Tuition_Fee_Year2", "Bus_Fee_Year2", "Total_Year2",
    "Tuition_Fee_Year3", "Bus_Fee_Year3", "Total_Year3",
    "Tuition_Fee_Year4", "Bus_Fee_Year4", "Total_Year4",
    "Grand_Total"
]

def ensure_excel_dir():
    if not os.path.exists(EXCEL_DIR):
        os.makedirs(EXCEL_DIR)

def sanitize_filename(name: str) -> str:
    clean = re.sub(r'[^a-zA-Z0-9]', '_', name.strip())
    return re.sub(r'_+', '_', clean).strip('_')

def get_course_code(course_name: str) -> str:
    text = re.sub(r'[^a-zA-Z\s]', '', course_name).strip()
    words = [w.upper() for w in text.split() if w.lower() not in ['b', 'e', 'tech', 'and', 'in', 'of']]
    if words:
        return "".join(w[0] for w in words)
    return "CEC"

def format_serial_no(raw_sno, course_name: str) -> str:
    code = get_course_code(course_name)
    sno_str = str(raw_sno).strip() or "1"
    return f"{sno_str}/{code}/CAHCET"

def parse_year_sheet_name(class_str: str) -> str:
    cl_lower = str(class_str).lower()
    if "1" in cl_lower or "first" in cl_lower:
        return "1_Year"
    elif "2" in cl_lower or "second" in cl_lower:
        return "2_Year"
    elif "3" in cl_lower or "third" in cl_lower:
        return "3_Year"
    elif "4" in cl_lower or "fourth" in cl_lower:
        return "4_Year"
    return "General"

def process_academic_year(acad_input: str) -> str:
    val = str(acad_input or "").strip()

    # Rule 1 & 2: Parse short format (e.g., 25-26) or standard format (e.g., 2025-2026)
    match_short = re.match(r'^(\d{2})[-/](\d{2})$', val)
    match_long = re.match(r'^(\d{4})[-/](\d{4})$', val)

    if match_short:
        y1, y2 = match_short.groups()
        return f"20{y1} – 20{y2}"
    elif match_long:
        y1, y2 = match_long.groups()
        return f"{y1} – {y2}"
    elif val:
        return val

    # Rule 3: Dynamic fallback based on date threshold (June to May academic cycle)
    today = date.today()
    current_year = today.year
    current_month = today.month

    # If date is June or later in current year, academic year is YYYY – (YYYY+1)
    # If date is before June, academic year is (YYYY-1) – YYYY
    if current_month >= 6:
        start_yr = current_year
        end_yr = current_year + 1
    else:
        start_yr = current_year - 1
        end_yr = current_year

    return f"{start_yr} – {end_yr}"

def get_course_filepath(course_name: str) -> str:
    ensure_excel_dir()
    filename = f"{sanitize_filename(course_name)}.xlsx"
    return os.path.join(EXCEL_DIR, filename)

def init_course_db(course_name: str, class_str: str):
    ensure_excel_dir()
    filepath = get_course_filepath(course_name)
    sheet_name = parse_year_sheet_name(class_str)

    if not os.path.exists(filepath):
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        ws.append(COLUMNS)
        wb.save(filepath)
    else:
        wb = load_workbook(filepath)
        if sheet_name not in wb.sheetnames:
            ws = wb.create_sheet(title=sheet_name)
            ws.append(COLUMNS)
            wb.save(filepath)
        wb.close()

def get_all_course_files():
    ensure_excel_dir()
    files = []
    for f in os.listdir(EXCEL_DIR):
        if f.endswith(".xlsx") and not f.startswith("~$"):
            files.append(os.path.join(EXCEL_DIR, f))
    return files

def calculate_row_totals(data: dict) -> dict:
    processed = {}
    
    processed["S_No"] = str(data.get("S_No", "")).strip() or "1"
    processed["Date"] = date.today().strftime("%d-%m-%Y")
    processed["Roll_No"] = str(data.get("Roll_No", "")).strip()
    processed["Register_No"] = str(data.get("Register_No", "")).strip()
    processed["Student_Name"] = str(data.get("Student_Name", "")).strip()
    processed["Father_Name"] = str(data.get("Father_Name", "")).strip()
    processed["Class"] = str(data.get("Class", "")).strip()
    processed["Course"] = str(data.get("Course", "")).strip()
    processed["Academic_Year"] = process_academic_year(data.get("Academic_Year", ""))

    grand_total = 0.0
    for i in range(1, 5):
        t_key, b_key = f"Tuition_Fee_Year{i}", f"Bus_Fee_Year{i}"
        
        try:
            t_val = float(data.get(t_key) or 0)
        except ValueError:
            raise ValueError(f"Invalid Tuition Fee Year {i}. Enter numeric values only.")
            
        try:
            b_val = float(data.get(b_key) or 0)
        except ValueError:
            raise ValueError(f"Invalid Bus Fee Year {i}. Enter numeric values only.")
        
        year_total = t_val + b_val
        grand_total += year_total
        
        processed[t_key] = t_val
        processed[b_key] = b_val
        processed[f"Total_Year{i}"] = year_total

    processed["Grand_Total"] = grand_total
    return processed

def sort_sheet_by_roll_no(ws):
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) <= 1:
        return
    
    header = rows[0]
    data_rows = rows[1:]

    def roll_sort_key(row):
        val = str(row[2]).strip() if row[2] is not None else ""
        return int(val) if val.isdigit() else val

    data_rows.sort(key=roll_sort_key)

    ws.delete_rows(1, ws.max_row)
    ws.append(header)
    for index, r in enumerate(data_rows, start=1):
        r_list = list(r)
        r_list[0] = index
        ws.append(r_list)

def add_student_record(student_data: dict) -> bool:
    processed = calculate_row_totals(student_data)
    course = processed["Course"]
    class_str = processed["Class"]
    
    init_course_db(course, class_str)
    filepath = get_course_filepath(course)
    sheet_name = parse_year_sheet_name(class_str)
    
    existing = find_student_by_roll(processed["Roll_No"])
    if existing:
        return False

    wb = load_workbook(filepath)
    ws = wb[sheet_name]
    
    row_to_add = [processed.get(col) for col in COLUMNS]
    ws.append(row_to_add)
    sort_sheet_by_roll_no(ws)
    
    wb.save(filepath)
    return True

def update_student_record(student_data: dict) -> bool:
    processed = calculate_row_totals(student_data)
    target_roll = processed["Roll_No"]
    
    delete_student_by_roll(target_roll)
    return add_student_record(student_data)

def delete_student_by_roll(roll_no: str) -> bool:
    target_roll = str(roll_no).strip()
    deleted = False

    for filepath in get_all_course_files():
        try:
            wb = load_workbook(filepath)
            for s_name in wb.sheetnames:
                ws = wb[s_name]
                for r_idx in range(2, ws.max_row + 1):
                    cell_val = str(ws.cell(row=r_idx, column=3).value).strip()
                    if cell_val == target_roll:
                        ws.delete_rows(r_idx, 1)
                        sort_sheet_by_roll_no(ws)
                        deleted = True
                        break
                if deleted:
                    break
            if deleted:
                wb.save(filepath)
                break
            wb.close()
        except Exception:
            continue
            
    return deleted

def format_student_display_fields(student_dict: dict) -> dict:
    if not student_dict:
        return student_dict

    # 1. Student Name -> UPPERCASE
    if student_dict.get("Student_Name"):
        student_dict["Student_Name"] = str(student_dict["Student_Name"]).strip().upper()

    # 2. Father Name -> Title Case & prefix "Mr."
    if student_dict.get("Father_Name"):
        f_name = str(student_dict["Father_Name"]).strip()
        f_name_title = f_name.title()
        if not f_name_title.startswith("Mr.") and not f_name_title.startswith("Mr "):
            student_dict["Father_Name"] = f"Mr. {f_name_title}"
        else:
            student_dict["Father_Name"] = f_name_title
    else:
        student_dict["Father_Name"] = ""

    # 3. Class -> Title Case
    if student_dict.get("Class"):
        student_dict["Class"] = str(student_dict["Class"]).strip().title()

    # 4. Course -> Title Case
    if student_dict.get("Course"):
        student_dict["Course"] = str(student_dict["Course"]).strip().title()

    # 5. Academic Year Fallback Formatting
    student_dict["Academic_Year"] = process_academic_year(student_dict.get("Academic_Year", ""))

    return student_dict

def find_student_by_roll(roll_no: str) -> dict:
    target_roll = str(roll_no).strip()

    for filepath in get_all_course_files():
        try:
            wb = load_workbook(filepath)
            for s_name in wb.sheetnames:
                ws = wb[s_name]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if row[2] is not None and str(row[2]).strip() == target_roll:
                        student_dict = {}
                        for idx, col_name in enumerate(COLUMNS):
                            val = row[idx] if idx < len(row) else ""
                            if val is None or val == "":
                                student_dict[col_name] = 0 if ("Fee" in col_name or "Total" in col_name) else ""
                            else:
                                student_dict[col_name] = val
                        
                        student_dict["Formatted_SNo"] = format_serial_no(student_dict["S_No"], student_dict["Course"])
                        wb.close()
                        return format_student_display_fields(student_dict)
            wb.close()
        except Exception:
            continue

    return None